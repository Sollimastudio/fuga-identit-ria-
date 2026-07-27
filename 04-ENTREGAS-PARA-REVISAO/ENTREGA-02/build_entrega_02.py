#!/usr/bin/env python3
"""Build Entrega 2: diagrams plus a reflowable 14 x 21 cm DOCX.

The manuscript remains the source of truth. This script generates the four
functional diagrams as editable SVG and high-resolution PNG, then places them
inline in a semantically styled DOCX suitable for revision and Kindle Create.
"""

from __future__ import annotations

import argparse
import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_entrega_00 import (
    BODY_FONT,
    DISPLAY_FONT,
    INK,
    IVORY,
    MUTED,
    NAVY,
    SANS_FONT,
    TEAL,
    VERMILION,
    add_inline_markdown,
    add_page_number,
    add_paragraph_left_border,
    configure_section,
    configure_styles,
    set_alt_text,
    set_keep_lines,
    set_page_number_start,
    set_paragraph_shading,
    set_run_font,
    set_style_font,
)


COMMANDS = (
    "IDENTIFIQUE A NARRATIVA.",
    "SEPARE FATO DE INTERPRETAÇÃO.",
    "LOCALIZE O OMITIDO.",
    "DESCUBRA O REBATIZADO.",
    "IDENTIFIQUE O VALOR OPERANTE.",
    "NOMEIE A CONTRADIÇÃO.",
    "LOCALIZE OS POLOS.",
    "ENCONTRE O EIXO.",
    "TESTE A LÓGICA.",
    "TESTE A SIMETRIA.",
)

FIGURES = {
    "FIG-03": {
        "filename": "FIG-03_DO_ACONTECIMENTO_AO_COMANDO.png",
        "title": "Do acontecimento ao comando",
        "caption": "FIGURA 3  •  Do acontecimento ao comando",
        "alt": (
            "Fluxo vertical em cinco etapas. Acontecimento ou fato pergunta o "
            "que ocorreu; recorte pergunta o que entrou e saiu; interpretação "
            "pergunta qual sentido foi atribuído; reclassificação pergunta qual "
            "nome moral mudou; comando pergunta o que se deve pensar, sentir ou "
            "fazer. Cada etapa acrescenta algo à anterior."
        ),
    },
    "FIG-04": {
        "filename": "FIG-04_MATRIZ_NARRATIVA.png",
        "title": "Matriz NARRATIVA",
        "caption": "FIGURA 4  •  Matriz NARRATIVA: nove movimentos de leitura",
        "alt": (
            "Nove linhas verticais apresentam a Matriz NARRATIVA: Núcleo, "
            "Autoridade, Recorte, Reclassificação, Afeto e atores, Testabilidade, "
            "Imunização, Valores e Ação e autoria. Cada linha contém uma pergunta "
            "curta para examinar a narrativa."
        ),
    },
    "FIG-05": {
        "filename": "FIG-05_COMO_UMA_CRENSA_CHEGA_A_ACAO.png",
        "title": "Como uma crença chega à ação",
        "caption": "FIGURA 5  •  Como uma crença chega à ação",
        "alt": (
            "Fluxo vertical de crenças para hierarquia de valores, interpretação, "
            "decisão e consequência. Uma faixa lateral informa que emoção, "
            "contexto, poder, vínculos e recursos interferem em todo o percurso; "
            "o processo não é apresentado como determinista."
        ),
    },
    "FIG-06": {
        "filename": "FIG-06_POLOS_NAO_SAO_EIXO.png",
        "title": "Polos não são eixo",
        "caption": "FIGURA 6  •  Polos não são eixo: o critério decide a posição",
        "alt": (
            "Dois polos são submetidos ao mesmo critério verificável: realidade, "
            "valores, coerência, responsabilidade, consequências, simetria e "
            "reconhecimento do outro. O resultado é uma posição responsável, que "
            "pode coincidir com um polo quando os fatos justificarem."
        ),
    },
}


def svg_header(width: int, height: int, title: str, deck: str) -> list[str]:
    return [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}" role="img" aria-labelledby="title desc">',
        f"<title id=\"title\">{html.escape(title)}</title>",
        f"<desc id=\"desc\">{html.escape(deck)}</desc>",
        "<defs>",
        '<marker id="arrow" markerWidth="12" markerHeight="12" refX="9" refY="6" '
        'orient="auto" markerUnits="strokeWidth">',
        f'<path d="M0,0 L10,6 L0,12 z" fill="#{VERMILION}"/>',
        "</marker>",
        '<filter id="shadow" x="-10%" y="-10%" width="120%" height="130%">',
        '<feDropShadow dx="0" dy="5" stdDeviation="6" flood-color="#0B1020" flood-opacity=".10"/>',
        "</filter>",
        "</defs>",
        f'<rect width="{width}" height="{height}" fill="#{IVORY}"/>',
        f'<rect x="0" y="0" width="{width}" height="36" fill="#{VERMILION}"/>',
        f'<text x="90" y="118" font-family="{DISPLAY_FONT}" font-size="58" '
        f'font-weight="700" fill="#{NAVY}">{html.escape(title)}</text>',
        f'<text x="90" y="170" font-family="{SANS_FONT}" font-size="26" '
        f'fill="#{TEAL}">{html.escape(deck)}</text>',
        f'<line x1="90" y1="205" x2="{width - 90}" y2="205" '
        f'stroke="#{TEAL}" stroke-width="3"/>',
    ]


def svg_footer(width: int, height: int, note: str) -> list[str]:
    return [
        f'<text x="{width / 2}" y="{height - 55}" text-anchor="middle" '
        f'font-family="{SANS_FONT}" font-size="24" fill="#{MUTED}">'
        f"{html.escape(note)}</text>",
        "</svg>",
    ]


def diagram_03() -> str:
    width, height = 1200, 1450
    lines = svg_header(
        width,
        height,
        "Do acontecimento ao comando",
        "Cada camada acrescenta algo ao que ocorreu.",
    )
    stages = [
        ("1", "ACONTECIMENTO / FATO", "O que ocorreu?"),
        ("2", "RECORTE", "O que entrou e saiu?"),
        ("3", "INTERPRETAÇÃO", "Que sentido foi atribuído?"),
        ("4", "RECLASSIFICAÇÃO", "Que nome moral mudou?"),
        ("5", "COMANDO", "O que devo pensar ou fazer?"),
    ]
    y0, h, gap = 235, 175, 28
    for i, (num, name, question) in enumerate(stages):
        y = y0 + i * (h + gap)
        fill = "FFFFFF" if i % 2 == 0 else "F9F6F0"
        lines.extend(
            [
                f'<rect x="150" y="{y}" width="900" height="{h}" rx="24" '
                f'fill="#{fill}" stroke="#{TEAL}" stroke-width="3" filter="url(#shadow)"/>',
                f'<circle cx="235" cy="{y + 66}" r="42" '
                f'fill="#{VERMILION if i == 4 else NAVY}"/>',
                f'<text x="235" y="{y + 80}" text-anchor="middle" '
                f'font-family="{SANS_FONT}" font-size="40" font-weight="700" '
                f'fill="#FFFFFF">{num}</text>',
                f'<text x="320" y="{y + 67}" font-family="{SANS_FONT}" '
                f'font-size="34" font-weight="700" fill="#{NAVY}">{name}</text>',
                f'<text x="320" y="{y + 118}" font-family="{BODY_FONT}" '
                f'font-size="30" fill="#{INK}">{question}</text>',
            ]
        )
        if i < len(stages) - 1:
            y1 = y + h + 9
            y2 = y + h + gap - 10
            lines.append(
                f'<line x1="600" y1="{y1}" x2="600" y2="{y2}" '
                f'stroke="#{VERMILION}" stroke-width="6" marker-end="url(#arrow)"/>'
            )
    lines += svg_footer(
        width,
        height,
        "Volte pelo caminho inverso para testar se o comando decorre do fato.",
    )
    return "\n".join(lines)


def diagram_04() -> str:
    width, height = 1200, 1650
    lines = svg_header(
        width,
        height,
        "Matriz NARRATIVA",
        "Nove movimentos para transformar incômodo em perguntas.",
    )
    rows = [
        ("N", "NÚCLEO", "Qual tese sustenta o restante?"),
        ("A", "AUTORIDADE", "Por que confiar nesta voz?"),
        ("R", "RECORTE", "O que ficou de fora?"),
        ("R", "RECLASSIFICAÇÃO", "Que fato mudou de nome?"),
        ("A", "AFETO E ATORES", "O que devo sentir e ser?"),
        ("T", "TESTABILIDADE", "O que poderia corrigir a tese?"),
        ("I", "IMUNIZAÇÃO", "Como a crítica vira confirmação?"),
        ("V", "VALORES", "Qual valor vence com custo?"),
        ("A", "AÇÃO E AUTORIA", "O que devo fazer e entregar?"),
    ]
    y0, h, gap = 210, 128, 7
    for i, (letter, name, question) in enumerate(rows):
        y = y0 + i * (h + gap)
        fill = "FFFFFF" if i % 2 == 0 else "F9F6F0"
        accent = VERMILION if i in (3, 6, 8) else NAVY
        lines.extend(
            [
                f'<rect x="95" y="{y}" width="1010" height="{h}" rx="20" '
                f'fill="#{fill}" stroke="#{TEAL}" stroke-width="2"/>',
                f'<rect x="115" y="{y + 16}" width="112" height="96" rx="18" '
                f'fill="#{accent}"/>',
                f'<text x="171" y="{y + 82}" text-anchor="middle" '
                f'font-family="{SANS_FONT}" font-size="54" font-weight="700" '
                f'fill="#FFFFFF">{letter}</text>',
                f'<text x="270" y="{y + 54}" font-family="{SANS_FONT}" '
                f'font-size="30" font-weight="700" fill="#{NAVY}">{name}</text>',
                f'<text x="270" y="{y + 101}" font-family="{BODY_FONT}" '
                f'font-size="29" fill="#{INK}">{question}</text>',
            ]
        )
    lines += svg_footer(
        width,
        height,
        "A matriz produz perguntas verificáveis; não diagnostica pessoas ou grupos.",
    )
    return "\n".join(lines)


def diagram_05() -> str:
    width, height = 1200, 1650
    lines = svg_header(
        width,
        height,
        "Como uma crença chega à ação",
        "Um encadeamento influenciado, não um destino mecânico.",
    )
    stages = [
        ("CRENÇAS", "O que considero verdadeiro"),
        ("HIERARQUIA DE VALORES", "O que vence quando há custo"),
        ("INTERPRETAÇÃO", "O significado recebido"),
        ("DECISÃO", "A escolha possível no contexto"),
        ("CONSEQUÊNCIA", "O que reforça ou revisa"),
    ]
    y0, h, gap = 265, 180, 55
    for i, (name, detail) in enumerate(stages):
        y = y0 + i * (h + gap)
        lines.extend(
            [
                f'<rect x="105" y="{y}" width="660" height="{h}" rx="22" '
                f'fill="#FFFFFF" stroke="#{TEAL}" stroke-width="3" filter="url(#shadow)"/>',
                f'<rect x="105" y="{y}" width="18" height="{h}" rx="9" '
                f'fill="#{VERMILION if i in (1, 4) else NAVY}"/>',
                f'<text x="165" y="{y + 72}" font-family="{SANS_FONT}" '
                f'font-size="32" font-weight="700" fill="#{NAVY}">{name}</text>',
                f'<text x="165" y="{y + 125}" font-family="{BODY_FONT}" '
                f'font-size="27" fill="#{INK}">{detail}</text>',
            ]
        )
        if i < len(stages) - 1:
            lines.append(
                f'<line x1="435" y1="{y + h + 10}" x2="435" '
                f'y2="{y + h + gap - 12}" stroke="#{VERMILION}" '
                f'stroke-width="6" marker-end="url(#arrow)"/>'
            )
    lines.extend(
        [
            f'<rect x="820" y="365" width="285" height="880" rx="26" '
            f'fill="#{NAVY}"/>',
            f'<text x="962" y="435" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="29" font-weight="700" fill="#FFFFFF">INTERFEREM</text>',
            f'<text x="962" y="475" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="25" fill="#F4EFE6">em todo o percurso</text>',
        ]
    )
    factors = ["EMOÇÃO", "CONTEXTO", "PODER", "VÍNCULOS", "RECURSOS"]
    for i, factor in enumerate(factors):
        y = 535 + i * 128
        lines.extend(
            [
                f'<rect x="865" y="{y}" width="195" height="84" rx="16" '
                f'fill="#{TEAL}"/>',
                f'<text x="962" y="{y + 53}" text-anchor="middle" '
                f'font-family="{SANS_FONT}" font-size="25" font-weight="700" '
                f'fill="#FFFFFF">{factor}</text>',
            ]
        )
    lines += svg_footer(
        width,
        height,
        "Autoria enxerga onde há escolha e onde a escolha foi reduzida.",
    )
    return "\n".join(lines)


def diagram_06() -> str:
    width, height = 1200, 1700
    lines = svg_header(
        width,
        height,
        "Polos não são eixo",
        "A posição nasce do critério, não da distância entre extremos.",
    )
    lines.extend(
        [
            f'<rect x="105" y="270" width="390" height="175" rx="24" '
            f'fill="#{NAVY}"/>',
            f'<text x="300" y="342" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="36" font-weight="700" fill="#FFFFFF">POLO A</text>',
            f'<text x="300" y="393" text-anchor="middle" font-family="{BODY_FONT}" '
            f'font-size="28" fill="#F4EFE6">Resposta absolutizada</text>',
            f'<rect x="705" y="270" width="390" height="175" rx="24" '
            f'fill="#{VERMILION}"/>',
            f'<text x="900" y="342" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="36" font-weight="700" fill="#FFFFFF">POLO B</text>',
            f'<text x="900" y="393" text-anchor="middle" font-family="{BODY_FONT}" '
            f'font-size="28" fill="#FFFFFF">Resposta absolutizada</text>',
            f'<line x1="300" y1="455" x2="500" y2="550" stroke="#{TEAL}" '
            f'stroke-width="7" marker-end="url(#arrow)"/>',
            f'<line x1="900" y1="455" x2="700" y2="550" stroke="#{TEAL}" '
            f'stroke-width="7" marker-end="url(#arrow)"/>',
            f'<rect x="275" y="570" width="650" height="135" rx="25" '
            f'fill="#FFFFFF" stroke="#{VERMILION}" stroke-width="5" filter="url(#shadow)"/>',
            f'<text x="600" y="625" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="31" font-weight="700" fill="#{NAVY}">EIXO = CRITÉRIO VERIFICÁVEL</text>',
            f'<text x="600" y="668" text-anchor="middle" font-family="{BODY_FONT}" '
            f'font-size="27" fill="#{INK}">O mesmo teste para ambos os polos</text>',
        ]
    )
    criteria = [
        "REALIDADE",
        "VALORES",
        "COERÊNCIA",
        "RESPONSABILIDADE",
        "CONSEQUÊNCIAS",
        "SIMETRIA",
        "RECONHECIMENTO DO OUTRO",
    ]
    positions = [
        (110, 770, 310),
        (445, 770, 310),
        (780, 770, 310),
        (110, 890, 470),
        (605, 890, 485),
        (110, 1010, 310),
        (445, 1010, 645),
    ]
    for criterion, (x, y, w) in zip(criteria, positions):
        lines.extend(
            [
                f'<rect x="{x}" y="{y}" width="{w}" height="85" rx="17" '
                f'fill="#{TEAL}"/>',
                f'<text x="{x + w / 2}" y="{y + 54}" text-anchor="middle" '
                f'font-family="{SANS_FONT}" font-size="24" font-weight="700" '
                f'fill="#FFFFFF">{criterion}</text>',
            ]
        )
    lines.extend(
        [
            f'<line x1="600" y1="1128" x2="600" y2="1218" '
            f'stroke="#{VERMILION}" stroke-width="7" marker-end="url(#arrow)"/>',
            f'<rect x="200" y="1240" width="800" height="190" rx="28" '
            f'fill="#{NAVY}" stroke="#{NAVY}" stroke-width="3"/>',
            f'<text x="600" y="1315" text-anchor="middle" font-family="{SANS_FONT}" '
            f'font-size="40" font-weight="700" fill="#FFFFFF">POSIÇÃO RESPONSÁVEL</text>',
            f'<text x="600" y="1373" text-anchor="middle" font-family="{BODY_FONT}" '
            f'font-size="29" fill="#F4EFE6">Pode coincidir com um polo</text>',
        ]
    )
    lines += svg_footer(
        width,
        height,
        "Eixo não é ponto médio: é o lugar que os fatos e os valores sustentam.",
    )
    return "\n".join(lines)


def generate_diagrams(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    diagrams = {
        "FIG-03_DO_ACONTECIMENTO_AO_COMANDO": diagram_03(),
        "FIG-04_MATRIZ_NARRATIVA": diagram_04(),
        "FIG-05_COMO_UMA_CRENSA_CHEGA_A_ACAO": diagram_05(),
        "FIG-06_POLOS_NAO_SAO_EIXO": diagram_06(),
    }
    for stem, svg in diagrams.items():
        svg_path = output_dir / f"{stem}.svg"
        png_path = output_dir / f"{stem}.png"
        svg_path.write_text(svg, encoding="utf-8")
        subprocess.run(
            [
                "inkscape",
                str(svg_path),
                "--export-type=png",
                f"--export-filename={png_path}",
                "--export-width=2000",
                "--export-background-opacity=0",
            ],
            check=True,
            capture_output=True,
            text=True,
        )


def configure_part_styles(doc: Document) -> None:
    configure_styles(doc)
    styles = doc.styles

    for name in (
        "Part Title",
        "Part Deck",
        "Part Body",
        "Chapter Kicker",
        "References",
    ):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    set_style_font(styles["Part Title"], DISPLAY_FONT, 24, NAVY, True)
    styles["Part Title"].paragraph_format.space_before = Pt(2)
    styles["Part Title"].paragraph_format.space_after = Pt(17)
    styles["Part Title"].paragraph_format.keep_with_next = True

    set_style_font(styles["Part Deck"], BODY_FONT, 12.2, NAVY)
    styles["Part Deck"].font.italic = True
    styles["Part Deck"].paragraph_format.line_spacing = 1.16
    styles["Part Deck"].paragraph_format.space_after = Pt(8)

    set_style_font(styles["Part Body"], BODY_FONT, 10.7, INK)
    styles["Part Body"].paragraph_format.line_spacing = 1.30
    styles["Part Body"].paragraph_format.space_after = Pt(6.5)
    styles["Part Body"].paragraph_format.widow_control = True

    set_style_font(styles["Chapter Kicker"], SANS_FONT, 8.4, VERMILION, True)
    styles["Chapter Kicker"].paragraph_format.space_after = Pt(7)
    styles["Chapter Kicker"].paragraph_format.keep_with_next = True

    set_style_font(styles["References"], BODY_FONT, 6.8, INK)
    styles["References"].paragraph_format.left_indent = Cm(0.25)
    styles["References"].paragraph_format.first_line_indent = Cm(-0.25)
    styles["References"].paragraph_format.line_spacing = 1.0
    styles["References"].paragraph_format.space_after = Pt(0.5)
    styles["References"].paragraph_format.widow_control = True

    styles["Heading 1"].font.size = Pt(22)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].paragraph_format.space_before = Pt(5)
    styles["Heading 2"].paragraph_format.space_after = Pt(10)

    h3 = styles["Heading 3"]
    set_style_font(h3, SANS_FONT, 10.5, NAVY, True)
    h3.paragraph_format.space_before = Pt(7)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    styles["Book Lead"].font.size = Pt(12.6)
    styles["Book Lead"].paragraph_format.space_after = Pt(8)
    styles["Command Box"].font.bold = False

    for style_name in ("List Number", "List Bullet"):
        style = styles[style_name]
        set_style_font(style, BODY_FONT, 10.1, INK)
        style.paragraph_format.left_indent = Cm(0.62)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.line_spacing = 1.22
        style.paragraph_format.space_after = Pt(4)


def enable_mirrored_margins(doc: Document) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:mirrorMargins")) is None:
        settings.append(OxmlElement("w:mirrorMargins"))


def set_two_columns(section, space_twips: int = 260) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(space_twips))


def add_delivery_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(30)
    add_paragraph_left_border(p, VERMILION, size=28, space=12)
    r = p.add_run("FUGA IDENTITÁRIA")
    set_run_font(r, SANS_FONT, 8.8, TEAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("PARTE II")
    set_run_font(r, SANS_FONT, 9.5, VERMILION, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(15)
    r = p.add_run("LEIA O QUE\nLÊ VOCÊ")
    set_run_font(r, DISPLAY_FONT, 31, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(33)
    add_paragraph_left_border(p, TEAL, size=12, space=9)
    r = p.add_run(
        "Narrativa • Matriz NARRATIVA • crenças e valores • lógica da situação"
    )
    set_run_font(r, BODY_FONT, 10.2, INK, italic=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(25)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SOL LIMA")
    set_run_font(r, SANS_FONT, 11.5, NAVY, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("Entrega 2 — manuscrito editorial para revisão")
    set_run_font(r, BODY_FONT, 8.8, MUTED, italic=True)


def add_running_furniture(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("FUGA IDENTITÁRIA")
    set_run_font(r, SANS_FONT, 7.5, NAVY, bold=True)
    r = p.add_run("   •   PARTE II")
    set_run_font(r, SANS_FONT, 7.5, MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("SOL LIMA   •   ")
    set_run_font(r, SANS_FONT, 7.2, MUTED)
    add_page_number(p)
    for run in p.runs:
        set_run_font(run, SANS_FONT, 7.2, MUTED)


def add_chapter_heading(doc: Document, title: str) -> None:
    match = re.match(r"^(\d+)\.\s+(.+)$", title)
    if not match:
        p = doc.add_paragraph(style="Heading 1")
        r = p.add_run(title)
        set_run_font(r, DISPLAY_FONT, 22, NAVY, bold=True)
        return
    number, chapter_title = match.groups()
    kicker = doc.add_paragraph(style="Chapter Kicker")
    kicker.paragraph_format.page_break_before = True
    r = kicker.add_run(f"CAPÍTULO {number}")
    set_run_font(r, SANS_FONT, 8.4, VERMILION, bold=True)

    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = False
    r = p.add_run(chapter_title)
    set_run_font(r, DISPLAY_FONT, 22, NAVY, bold=True)


def add_blockquote(doc: Document, block: list[str]) -> None:
    text = "\n".join(line for line in block if line)
    is_command = any(command in text for command in COMMANDS)
    style = "Command Box" if is_command else "Editorial Quote"
    p = doc.add_paragraph(style=style)
    add_paragraph_left_border(p, VERMILION, size=18 if is_command else 12, space=8)
    set_paragraph_shading(p, IVORY)
    add_inline_markdown(
        p,
        text,
        base_font=SANS_FONT if is_command else BODY_FONT,
        base_size=10.1 if is_command else 11.1,
    )
    set_keep_lines(p)


def add_figure(doc: Document, figure_id: str, figures_dir: Path) -> None:
    data = FIGURES[figure_id]
    figure_path = figures_dir / data["filename"]
    figure_width = 8.4 if figure_id == "FIG-04" else 9.6
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(figure_path), width=Cm(figure_width))
    set_alt_text(shape, data["title"], data["alt"])

    cap = doc.add_paragraph(style="Figure Caption")
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run(data["caption"])
    set_run_font(r, SANS_FONT, 7.5, MUTED, italic=True)


def build_document(source: Path, figures_dir: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    start = lines.index("# PARTE II")
    lines = lines[start:]

    doc = Document()
    configure_part_styles(doc)
    enable_mirrored_margins(doc)
    configure_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    add_delivery_title_page(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    body_section.different_first_page_header_footer = False
    set_page_number_start(body_section, 61)
    add_running_furniture(body_section)

    props = doc.core_properties
    props.title = "Fuga Identitária — Entrega 2 — Parte II"
    props.author = "Sol Lima"
    props.subject = "Leia o que lê você"
    props.keywords = (
        "fuga identitária, narrativa, Matriz NARRATIVA, crenças, valores, "
        "metacognição, lógica, autoria"
    )
    props.comments = "Versão editorial da Entrega 2 para revisão da autora."

    part_opening = True
    part_paragraphs = 0
    chapter_paragraphs = 0
    in_references = False
    reference_intro_added = False
    reference_columns_started = False
    blockquote: list[str] = []

    def flush_blockquote() -> None:
        nonlocal blockquote
        if blockquote:
            add_blockquote(doc, blockquote)
            blockquote = []

    for raw in lines:
        stripped = raw.strip()

        if stripped.startswith(">"):
            blockquote.append(stripped[1:].lstrip())
            continue
        flush_blockquote()

        if not stripped or stripped == "---":
            continue

        figure_match = re.match(r"!\[(FIG-\d+)[^\]]*\]\([^)]+\)", stripped)
        if figure_match:
            add_figure(doc, figure_match.group(1), figures_dir)
            continue

        if stripped == "# PARTE II":
            continue

        if stripped.startswith("# "):
            title = stripped[2:].strip()
            if title == "Notas de pesquisa da Entrega 2":
                in_references = True
                part_opening = False
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.page_break_before = True
                p.paragraph_format.space_after = Pt(7)
                r = p.add_run(title)
                set_run_font(r, DISPLAY_FONT, 18, NAVY, bold=True)
            else:
                part_opening = False
                chapter_paragraphs = 0
                add_chapter_heading(doc, title)
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, stripped[4:].strip(), base_font=SANS_FONT, base_size=10.5)
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if part_opening:
                p = doc.add_paragraph(style="Part Title")
                r = p.add_run(title)
                set_run_font(r, DISPLAY_FONT, 24, NAVY, bold=True)
            else:
                p = doc.add_paragraph(style="Heading 2")
                add_inline_markdown(p, title, base_font=BODY_FONT, base_size=13)
            continue

        if in_references:
            if not reference_intro_added:
                p = doc.add_paragraph(style="Normal")
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.10
                add_inline_markdown(p, stripped, base_size=8.2)
                reference_intro_added = True
                continue
            if not reference_columns_started:
                refs_section = doc.add_section(WD_SECTION.CONTINUOUS)
                configure_section(refs_section)
                set_two_columns(refs_section)
                reference_columns_started = True
            p = doc.add_paragraph(style="References")
            add_inline_markdown(p, stripped, base_size=6.8)
            set_keep_lines(p)
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, ordered.group(2), base_size=10.1)
            set_keep_lines(p)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1), base_size=10.1)
            set_keep_lines(p)
            continue

        if part_opening:
            style = "Part Deck" if part_paragraphs == 0 else "Part Body"
            p = doc.add_paragraph(style=style)
            add_inline_markdown(
                p,
                stripped,
                base_size=12.2 if part_paragraphs == 0 else 10.7,
            )
            set_keep_lines(p)
            part_paragraphs += 1
            continue

        if chapter_paragraphs == 0:
            p = doc.add_paragraph(style="Book Lead")
            add_inline_markdown(p, stripped, base_size=12.6)
        else:
            p = doc.add_paragraph(style="Normal")
            add_inline_markdown(p, stripped)
            set_keep_lines(p)
        chapter_paragraphs += 1

    flush_blockquote()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("figures_dir", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    generate_diagrams(args.figures_dir)
    build_document(args.source, args.figures_dir, args.output)


if __name__ == "__main__":
    main()
