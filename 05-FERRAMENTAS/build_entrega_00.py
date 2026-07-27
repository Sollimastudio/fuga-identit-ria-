#!/usr/bin/env python3
"""Build Entrega 0 as a polished 14 x 21 cm DOCX for review and later fusion."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = "0B1020"
VERMILION = "D84A3A"
IVORY = "F4EFE6"
TEAL = "5F7F83"
INK = "20242B"
MUTED = "666B73"
LIGHT_RULE = "D9D4CA"
WHITE = "FFFFFF"

BODY_FONT = "Bitstream Charter"
DISPLAY_FONT = "Nimbus Sans Narrow"
SANS_FONT = "Nimbus Sans"

PAGE_WIDTH_CM = 14.0
PAGE_HEIGHT_CM = 21.0
MARGIN_TOP_CM = 1.6
MARGIN_BOTTOM_CM = 1.75
MARGIN_LEFT_CM = 1.8
MARGIN_RIGHT_CM = 1.5
CONTENT_WIDTH_CM = PAGE_WIDTH_CM - MARGIN_LEFT_CM - MARGIN_RIGHT_CM


def set_run_font(
    run,
    name: str,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name: str, size: float, color: str, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def configure_section(section) -> None:
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(MARGIN_TOP_CM)
    section.bottom_margin = Cm(MARGIN_BOTTOM_CM)
    section.left_margin = Cm(MARGIN_LEFT_CM)
    section.right_margin = Cm(MARGIN_RIGHT_CM)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)
    section.gutter = Cm(0)


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, cached, end])


def add_paragraph_left_border(paragraph, color: str, size: int = 18, space: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_keep_lines(paragraph, keep_with_next: bool = False) -> None:
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.widow_control = True


def add_inline_markdown(paragraph, text: str, *, base_font: str = BODY_FONT, base_size: float = 10.4) -> None:
    token_re = re.compile(r"(\*\*.*?\*\*|\*[^*].*?\*)", re.DOTALL)
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            plain = text[cursor : match.start()]
            for idx, segment in enumerate(plain.split("\n")):
                if idx:
                    paragraph.add_run().add_break()
                if segment:
                    run = paragraph.add_run(segment)
                    set_run_font(run, base_font, base_size, INK)
        token = match.group(0)
        if token.startswith("**"):
            content = token[2:-2]
            bold, italic = True, False
        else:
            content = token[1:-1]
            bold, italic = False, True
        for idx, segment in enumerate(content.split("\n")):
            if idx:
                paragraph.add_run().add_break()
            if segment:
                run = paragraph.add_run(segment)
                set_run_font(run, base_font, base_size, INK, bold=bold, italic=italic)
        cursor = match.end()
    if cursor < len(text):
        tail = text[cursor:]
        for idx, segment in enumerate(tail.split("\n")):
            if idx:
                paragraph.add_run().add_break()
            if segment:
                run = paragraph.add_run(segment)
                set_run_font(run, base_font, base_size, INK)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_paragraph_left_border(p, VERMILION, size=28, space=12)
    r = p.add_run("ENSAIO AUTORAL INVESTIGATIVO")
    set_run_font(r, SANS_FONT, 8.6, TEAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    set_keep_lines(p, True)
    r = p.add_run("FUGA")
    set_run_font(r, DISPLAY_FONT, 28, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    set_keep_lines(p, True)
    r = p.add_run("IDENTITÁRIA")
    set_run_font(r, DISPLAY_FONT, 35, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("A ANATOMIA DA DILUIÇÃO DO EU")
    set_run_font(r, SANS_FONT, 13, VERMILION, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SOL LIMA")
    set_run_font(r, SANS_FONT, 12, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Entrega 0 — Nota da autora, nota conceitual e introdução")
    set_run_font(r, BODY_FONT, 9.3, MUTED, italic=True)


def add_running_furniture(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("FUGA IDENTITÁRIA")
    set_run_font(r, SANS_FONT, 7.5, NAVY, bold=True)
    r = p.add_run("   •   ABERTURA")
    set_run_font(r, SANS_FONT, 7.5, MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("SOL LIMA   •   ")
    set_run_font(r, SANS_FONT, 7.2, MUTED)
    add_page_number(p)
    for run in p.runs:
        set_run_font(run, SANS_FONT, 7.2, MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, BODY_FONT, 10.4, INK)
    # Ragged-right is the safest review/KDP source setting because it avoids
    # rivers of white in a narrow 14 x 21 cm measure when Word/LibreOffice
    # Portuguese hyphenation dictionaries differ. Print typesetting can
    # re-enable justified text with professional language-aware hyphenation.
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6.5)
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    set_style_font(h1, DISPLAY_FONT, 20, NAVY, True)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(15)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True

    h2 = styles["Heading 2"]
    set_style_font(h2, BODY_FONT, 13, TEAL, True)
    h2.paragraph_format.space_before = Pt(0)
    h2.paragraph_format.space_after = Pt(16)
    h2.paragraph_format.keep_with_next = True

    for name, style_type in (
        ("Book Lead", WD_STYLE_TYPE.PARAGRAPH),
        ("Index Part", WD_STYLE_TYPE.PARAGRAPH),
        ("Index Entry", WD_STYLE_TYPE.PARAGRAPH),
        ("Figure Caption", WD_STYLE_TYPE.PARAGRAPH),
        ("Editorial Quote", WD_STYLE_TYPE.PARAGRAPH),
        ("Command Box", WD_STYLE_TYPE.PARAGRAPH),
    ):
        if name not in styles:
            styles.add_style(name, style_type)

    lead = styles["Book Lead"]
    set_style_font(lead, BODY_FONT, 14, NAVY)
    lead.font.italic = True
    lead.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lead.paragraph_format.line_spacing = 1.15
    lead.paragraph_format.space_before = Pt(2)
    lead.paragraph_format.space_after = Pt(7)
    lead.paragraph_format.keep_with_next = True

    part = styles["Index Part"]
    set_style_font(part, SANS_FONT, 10.2, NAVY, True)
    part.paragraph_format.space_before = Pt(13)
    part.paragraph_format.space_after = Pt(4)
    part.paragraph_format.keep_with_next = True

    entry = styles["Index Entry"]
    set_style_font(entry, BODY_FONT, 8.9, INK)
    entry.paragraph_format.left_indent = Cm(0.55)
    entry.paragraph_format.first_line_indent = Cm(-0.55)
    entry.paragraph_format.space_before = Pt(0)
    entry.paragraph_format.space_after = Pt(3.2)
    entry.paragraph_format.line_spacing = 1.05
    entry.paragraph_format.widow_control = True

    caption = styles["Figure Caption"]
    set_style_font(caption, SANS_FONT, 7.5, MUTED)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(11)
    caption.paragraph_format.keep_together = True

    quote = styles["Editorial Quote"]
    set_style_font(quote, BODY_FONT, 11.2, NAVY)
    quote.font.italic = True
    quote.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    quote.paragraph_format.left_indent = Cm(0.35)
    quote.paragraph_format.right_indent = Cm(0.2)
    quote.paragraph_format.space_before = Pt(8)
    quote.paragraph_format.space_after = Pt(10)
    quote.paragraph_format.line_spacing = 1.18
    quote.paragraph_format.keep_together = True

    command = styles["Command Box"]
    set_style_font(command, SANS_FONT, 11, NAVY, True)
    command.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    command.paragraph_format.left_indent = Cm(0.32)
    command.paragraph_format.right_indent = Cm(0.2)
    command.paragraph_format.space_before = Pt(10)
    command.paragraph_format.space_after = Pt(11)
    command.paragraph_format.line_spacing = 1.12
    command.paragraph_format.keep_together = True


def add_index_entry(doc: Document, line: str) -> None:
    match = re.match(r"^(\d+)\.\s+\*\*(.+?)\*\*\s+—\s+(.+)$", line)
    p = doc.add_paragraph(style="Index Entry")
    if not match:
        add_inline_markdown(p, line, base_size=8.9)
        return
    number, title, description = match.groups()
    r = p.add_run(f"{number}. ")
    set_run_font(r, SANS_FONT, 8.5, VERMILION, bold=True)
    r = p.add_run(title)
    set_run_font(r, BODY_FONT, 8.9, NAVY, bold=True)
    r = p.add_run(f" — {description}")
    set_run_font(r, BODY_FONT, 8.9, MUTED)


def add_blockquote(doc: Document, lines: list[str]) -> None:
    text = "\n".join(line for line in lines if line)
    is_command = "PARE." in text
    p = doc.add_paragraph(style="Command Box" if is_command else "Editorial Quote")
    add_paragraph_left_border(p, VERMILION, size=18 if is_command else 12, space=7)
    set_paragraph_shading(p, IVORY)
    add_inline_markdown(
        p,
        text,
        base_font=SANS_FONT if is_command else BODY_FONT,
        base_size=10.8 if is_command else 11.2,
    )
    set_keep_lines(p)


def add_figure(doc: Document, figure_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(figure_path), width=Cm(8.25))
    set_alt_text(
        shape,
        "O eu cercado por vozes fornecedoras",
        "Figura humana abstrata em espaço negativo, cercada por camadas e linhas que simbolizam fontes externas de identidade; uma linha vermelha vertical representa o eixo da autoria.",
    )
    cap = doc.add_paragraph(style="Figure Caption")
    r = cap.add_run("FIGURA 1  •  O eu cercado por vozes fornecedoras")
    set_run_font(r, SANS_FONT, 7.5, MUTED, italic=True)


def build_document(source: Path, figure_path: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    start = lines.index("# Sumário da travessia")
    lines = lines[start:]

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    add_title_page(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    set_page_number_start(body_section, 1)
    add_running_furniture(body_section)

    doc.core_properties.title = "Fuga Identitária — Entrega 0"
    doc.core_properties.author = "Sol Lima"
    doc.core_properties.subject = "Nota da autora, nota conceitual e introdução"
    doc.core_properties.keywords = "fuga identitária, autoria, narrativa, propósito, metacognição"
    doc.core_properties.comments = "Versão editorial para revisão da autora."

    in_index = False
    in_intro = False
    intro_body_count = 0
    blockquote: list[str] = []

    def flush_blockquote() -> None:
        nonlocal blockquote
        if blockquote:
            add_blockquote(doc, blockquote)
            blockquote = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith(">"):
            blockquote.append(stripped[1:].lstrip())
            continue
        flush_blockquote()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("![FIG-01"):
            add_figure(doc, figure_path)
            continue

        if stripped.startswith("# "):
            title = stripped[2:].strip()
            in_index = title == "Sumário da travessia"
            in_intro = title == "Introdução"
            intro_body_count = 0
            p = doc.add_paragraph(style="Heading 1")
            if title == "Sumário da travessia":
                p.paragraph_format.page_break_before = False
            r = p.add_run(title)
            set_run_font(r, DISPLAY_FONT, 20, NAVY, bold=True)
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if in_index:
                p = doc.add_paragraph(style="Index Part")
                r = p.add_run(title)
                set_run_font(r, SANS_FONT, 10.2, NAVY, bold=True)
            else:
                p = doc.add_paragraph(style="Heading 2")
                r = p.add_run(title)
                set_run_font(r, BODY_FONT, 13, TEAL, bold=True)
            continue

        if in_index and re.match(r"^\d+\.", stripped):
            add_index_entry(doc, stripped)
            continue

        if in_index and stripped.startswith("- "):
            p = doc.add_paragraph(style="Index Entry")
            add_inline_markdown(p, stripped[2:], base_size=8.9)
            continue

        if in_intro and intro_body_count < 2:
            p = doc.add_paragraph(style="Book Lead")
            add_inline_markdown(p, stripped, base_size=14)
            intro_body_count += 1
            continue

        p = doc.add_paragraph(style="Normal")
        add_inline_markdown(p, stripped)
        set_keep_lines(p)
        if in_intro:
            intro_body_count += 1

    flush_blockquote()

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("figure", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_document(args.source, args.figure, args.output)


if __name__ == "__main__":
    main()
