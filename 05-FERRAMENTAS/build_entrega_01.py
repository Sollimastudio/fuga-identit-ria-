#!/usr/bin/env python3
"""Build Entrega 1 as a polished 14 x 21 cm DOCX.

The module reuses the established Entrega 0 design system and adds Part I,
chapter-opening, references, command-box, and FIG-02 treatments.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_entrega_00 import (
    BODY_FONT,
    DISPLAY_FONT,
    INK,
    IVORY,
    MUTED,
    NAVY,
    SANS_FONT,
    TEAL,
    VERMILION,
    add_inline_markdown,
    add_page_number,
    add_paragraph_left_border,
    configure_section,
    configure_styles,
    set_alt_text,
    set_page_number_start,
    set_paragraph_shading,
    set_keep_lines,
    set_run_font,
    set_style_font,
)


COMMANDS = (
    "PARE.",
    "RECUPERE A PERGUNTA.",
    "SEPARE BUSCA DE FUGA.",
    "IDENTIFIQUE A FUGA.",
)


def configure_part_styles(doc: Document) -> None:
    configure_styles(doc)
    styles = doc.styles

    for name in (
        "Part Title",
        "Part Deck",
        "Part Body",
        "Chapter Kicker",
        "References",
    ):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    part_title = styles["Part Title"]
    set_style_font(part_title, DISPLAY_FONT, 22, NAVY, True)
    part_title.paragraph_format.space_before = Pt(4)
    part_title.paragraph_format.space_after = Pt(18)
    part_title.paragraph_format.keep_with_next = True

    part_deck = styles["Part Deck"]
    set_style_font(part_deck, BODY_FONT, 11.7, NAVY)
    part_deck.font.italic = True
    part_deck.paragraph_format.line_spacing = 1.12
    part_deck.paragraph_format.space_before = Pt(0)
    part_deck.paragraph_format.space_after = Pt(6)
    part_deck.paragraph_format.keep_with_next = False

    part_body = styles["Part Body"]
    set_style_font(part_body, BODY_FONT, 9.8, INK)
    part_body.paragraph_format.line_spacing = 1.16
    part_body.paragraph_format.space_before = Pt(0)
    part_body.paragraph_format.space_after = Pt(4.2)
    part_body.paragraph_format.widow_control = True

    chapter_kicker = styles["Chapter Kicker"]
    set_style_font(chapter_kicker, SANS_FONT, 8.2, VERMILION, True)
    chapter_kicker.paragraph_format.space_before = Pt(0)
    chapter_kicker.paragraph_format.space_after = Pt(7)
    chapter_kicker.paragraph_format.keep_with_next = True

    refs = styles["References"]
    set_style_font(refs, BODY_FONT, 7.0, INK)
    refs.paragraph_format.left_indent = Cm(0.24)
    refs.paragraph_format.first_line_indent = Cm(-0.24)
    refs.paragraph_format.line_spacing = 0.94
    refs.paragraph_format.space_before = Pt(0)
    refs.paragraph_format.space_after = Pt(1.8)
    refs.paragraph_format.widow_control = True

    styles["Heading 1"].font.size = Pt(22)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 2"].paragraph_format.space_before = Pt(2)
    styles["Heading 2"].paragraph_format.space_after = Pt(13)


def set_two_columns(section, space_twips: int = 260) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(space_twips))


def add_delivery_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(34)
    add_paragraph_left_border(p, VERMILION, size=28, space=12)
    r = p.add_run("FUGA IDENTITÁRIA")
    set_run_font(r, SANS_FONT, 8.8, TEAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("PARTE I")
    set_run_font(r, SANS_FONT, 9.5, VERMILION, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("ANTES DO\nPROPÓSITO,\nEXISTE UM EU")
    set_run_font(r, DISPLAY_FONT, 28, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(36)
    add_paragraph_left_border(p, TEAL, size=12, space=9)
    r = p.add_run(
        "A pergunta debaixo da pergunta • Buscar-se não é fugir de si • "
        "Duas fugas, o mesmo vazio"
    )
    set_run_font(r, BODY_FONT, 10.2, INK, italic=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SOL LIMA")
    set_run_font(r, SANS_FONT, 11.5, NAVY, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("Entrega 1 — manuscrito editorial para revisão")
    set_run_font(r, BODY_FONT, 8.8, MUTED, italic=True)


def add_running_furniture(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("FUGA IDENTITÁRIA")
    set_run_font(r, SANS_FONT, 7.5, NAVY, bold=True)
    r = p.add_run("   •   PARTE I")
    set_run_font(r, SANS_FONT, 7.5, MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("SOL LIMA   •   ")
    set_run_font(r, SANS_FONT, 7.2, MUTED)
    add_page_number(p)
    for run in p.runs:
        set_run_font(run, SANS_FONT, 7.2, MUTED)


def add_blockquote(doc: Document, lines: list[str]) -> None:
    text = "\n".join(line for line in lines if line)
    is_command = any(command in text for command in COMMANDS)
    style = "Command Box" if is_command else "Editorial Quote"
    p = doc.add_paragraph(style=style)
    add_paragraph_left_border(p, VERMILION, size=18 if is_command else 12, space=7)
    set_paragraph_shading(p, IVORY)
    add_inline_markdown(
        p,
        text,
        base_font=SANS_FONT if is_command else BODY_FONT,
        base_size=10.2 if is_command else 11.2,
    )
    set_keep_lines(p)


def add_figure(doc: Document, figure_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(figure_path), width=Cm(8.0))
    set_alt_text(
        shape,
        "Duas fugas, o mesmo ciclo",
        (
            "Figura humana abstrata diante de dois percursos circulares. "
            "À esquerda, ondas dispersas simbolizam a fuga cognitiva; à direita, "
            "máscaras e formas rígidas simbolizam a fuga identitária. Os caminhos "
            "convergem numa raiz escura, enquanto uma linha vermelha vertical "
            "representa a possibilidade de recuperar autoria."
        ),
    )
    cap = doc.add_paragraph(style="Figure Caption")
    r = cap.add_run(
        "FIGURA 2  •  Duas fugas, o mesmo ciclo de alívio e terceirização"
    )
    set_run_font(r, SANS_FONT, 7.5, MUTED, italic=True)


def add_chapter_heading(doc: Document, title: str) -> None:
    match = re.match(r"^(\d+)\.\s+(.+)$", title)
    if not match:
        p = doc.add_paragraph(style="Heading 1")
        r = p.add_run(title)
        set_run_font(r, DISPLAY_FONT, 22, NAVY, bold=True)
        return

    number, chapter_title = match.groups()
    kicker = doc.add_paragraph(style="Chapter Kicker")
    kicker.paragraph_format.page_break_before = True
    r = kicker.add_run(f"CAPÍTULO {number}")
    set_run_font(r, SANS_FONT, 8.2, VERMILION, bold=True)

    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = False
    r = p.add_run(chapter_title)
    set_run_font(r, DISPLAY_FONT, 22, NAVY, bold=True)


def build_document(source: Path, figure_path: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    start = lines.index("# PARTE I")
    lines = lines[start:]

    doc = Document()
    configure_part_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    add_delivery_title_page(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    set_page_number_start(body_section, 24)
    add_running_furniture(body_section)

    doc.core_properties.title = "Fuga Identitária — Entrega 1 — Parte I"
    doc.core_properties.author = "Sol Lima"
    doc.core_properties.subject = "Antes do propósito, existe um eu"
    doc.core_properties.keywords = (
        "fuga identitária, autoria, propósito, busca, fuga cognitiva, identidade"
    )
    doc.core_properties.comments = "Versão editorial da Entrega 1 para revisão da autora."

    in_references = False
    reference_intro_added = False
    reference_columns_started = False
    part_opening = True
    part_body_count = 0
    chapter_body_count = 0
    blockquote: list[str] = []

    def flush_blockquote() -> None:
        nonlocal blockquote
        if blockquote:
            add_blockquote(doc, blockquote)
            blockquote = []

    for raw in lines:
        stripped = raw.strip()

        if stripped.startswith(">"):
            blockquote.append(stripped[1:].lstrip())
            continue
        flush_blockquote()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("![FIG-02"):
            add_figure(doc, figure_path)
            continue

        if stripped == "# PARTE I":
            continue

        if stripped.startswith("# "):
            title = stripped[2:].strip()
            if title == "Notas de pesquisa da Entrega 1":
                in_references = True
                part_opening = False
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.space_after = Pt(7)
                r = p.add_run(title)
                set_run_font(r, DISPLAY_FONT, 18, NAVY, bold=True)
            else:
                part_opening = False
                chapter_body_count = 0
                add_chapter_heading(doc, title)
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if part_opening:
                p = doc.add_paragraph(style="Part Title")
                r = p.add_run(title)
                set_run_font(r, DISPLAY_FONT, 22, NAVY, bold=True)
            else:
                p = doc.add_paragraph(style="Heading 2")
                r = p.add_run(title)
                set_run_font(r, BODY_FONT, 13, TEAL, bold=True)
            continue

        if in_references:
            if not reference_intro_added:
                p = doc.add_paragraph(style="Normal")
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.12
                add_inline_markdown(p, stripped, base_size=8.2)
                reference_intro_added = True
                continue
            if not reference_columns_started:
                refs_section = doc.add_section(WD_SECTION.CONTINUOUS)
                configure_section(refs_section)
                set_two_columns(refs_section)
                reference_columns_started = True
            p = doc.add_paragraph(style="References")
            add_inline_markdown(p, stripped, base_size=7.0)
            set_keep_lines(p)
            continue

        if part_opening and part_body_count < 2:
            p = doc.add_paragraph(style="Part Deck")
            add_inline_markdown(p, stripped, base_size=11.7)
            part_body_count += 1
            continue

        if part_opening:
            p = doc.add_paragraph(style="Part Body")
            add_inline_markdown(p, stripped, base_size=9.8)
            set_keep_lines(p)
            part_body_count += 1
            continue

        if not part_opening and chapter_body_count < 2:
            p = doc.add_paragraph(style="Book Lead")
            add_inline_markdown(p, stripped, base_size=13.2)
            chapter_body_count += 1
            continue

        p = doc.add_paragraph(style="Normal")
        add_inline_markdown(p, stripped)
        set_keep_lines(p)
        chapter_body_count += 1

    flush_blockquote()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("figure", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_document(args.source, args.figure, args.output)


if __name__ == "__main__":
    main()
