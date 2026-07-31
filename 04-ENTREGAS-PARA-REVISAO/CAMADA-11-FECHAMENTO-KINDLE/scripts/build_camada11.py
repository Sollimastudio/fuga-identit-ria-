#!/usr/bin/env python3
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "FUGA_IDENTITARIA_CAMADA_11_FECHAMENTO_KINDLE.md"
OUTPUT = ROOT / "output"
FIGURES = ROOT / "figuras"
SCRIPTS = ROOT / "scripts"
REFERENCE = SCRIPTS / "reference_camada11_kindlereflow.docx"
COVER = FIGURES / "CAPA_CAMADA_11.png"
FONTS = ROOT / "fonts"
FONTCONFIG = FONTS / "fonts.conf"
DOCX_OUT = OUTPUT / "Fuga_Identitaria_Camada_11_Kindle_Reflow.docx"
PDF_OUT = OUTPUT / "Fuga_Identitaria_Camada_11_Prova_Tecnica.pdf"
EPUB_OUT = OUTPUT / "FUGA_IDENTITARIA_CAMADA_11_FECHAMENTO_KINDLE.epub"
COVER_JPG = OUTPUT / "CAPA_FUGA_IDENTITARIA_KINDLE_1600x2560.jpg"

TITLE = "Fuga Identitária"
SUBTITLE = "A Anatomia da Diluição do Eu"
AUTHOR = "Sol Lima"
DESCRIPTION = (
    "Uma investigação sobre como narrativas, pertencimentos e respostas prontas "
    "podem reduzir a autoria — e uma travessia prática para recuperar exame, "
    "responsabilidade e posicionamento."
)
SUBJECT = "Identidade; autoria; narrativas; pensamento crítico; autoconhecimento"
BUILD_DATE = "2026-07-31"

NAVY = "0B1020"
VERMILION = "D84A3A"
IVORY = "F4EFE6"
PETROL = "5F7F83"
INK = "20242B"


def run(
    cmd: list[str],
    cwd: Path | None = None,
    env: dict[str, str] | None = None,
) -> None:
    print("+", " ".join(cmd), flush=True)
    subprocess.run(cmd, cwd=str(cwd or ROOT), env=env, check=True)


def set_font(style, name: str, size: float | None = None, color: str | None = None) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)


def get_or_add_style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    for style in doc.styles:
        if style.name == name:
            return style
    return doc.styles.add_style(name, style_type)


def add_left_border(style, color: str, width: str = "18", space: str = "8") -> None:
    ppr = style._element.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), width)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    pbdr.append(left)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_obj = paragraph.add_run()
    run_obj.font.name = "Nimbus Sans"
    run_obj.font.size = Pt(8.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_obj._r.extend([begin, instr, separate, end])


def configure_section(section) -> None:
    section.page_width = Cm(14)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)
    section.different_first_page_header_footer = False


def add_mirror_margins(doc: Document) -> None:
    """Kindle source intentionally has no mirror-margin or page-field setting."""
    return None


def make_reference_docx() -> None:
    default_bytes = subprocess.check_output(
        ["pandoc", "--print-default-data-file", "reference.docx"]
    )
    REFERENCE.write_bytes(default_bytes)
    doc = Document(str(REFERENCE))
    for section in doc.sections:
        configure_section(section)

    normal = doc.styles["Normal"]
    set_font(normal, "Bitstream Charter", 10.4, INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    for style_name in ("Body Text", "First Paragraph", "Compact"):
        style = get_or_add_style(doc, style_name)
        style.base_style = normal
        set_font(style, "Bitstream Charter", 10.4, INK)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.30
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    set_font(title, "Nimbus Sans Narrow", 27, NAVY)
    title.font.bold = True
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(96)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    subtitle = get_or_add_style(doc, "Subtitle")
    set_font(subtitle, "Nimbus Sans Narrow", 15, PETROL)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    subtitle.paragraph_format.keep_with_next = True

    author = get_or_add_style(doc, "Author")
    set_font(author, "Nimbus Sans", 11, INK)
    author.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(0)

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 0, 18),
        ("Heading 2", 13.5, NAVY, 15, 7),
        ("Heading 3", 11.5, PETROL, 11, 5),
    ):
        style = get_or_add_style(doc, name)
        set_font(style, "Nimbus Sans Narrow", size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True

    quote = get_or_add_style(doc, "Block Text")
    quote.base_style = normal
    set_font(quote, "Bitstream Charter", 10.2, NAVY)
    quote.paragraph_format.left_indent = Cm(0.65)
    quote.paragraph_format.right_indent = Cm(0.25)
    quote.paragraph_format.space_before = Pt(7)
    quote.paragraph_format.space_after = Pt(7)
    quote.paragraph_format.keep_together = True
    add_left_border(quote, VERMILION)

    caption = get_or_add_style(doc, "Caption")
    set_font(caption, "Nimbus Sans", 8.4, PETROL)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    image_caption = get_or_add_style(doc, "Image Caption")
    image_caption.base_style = caption

    for style_name in ("Bullet List", "Ordered List"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        set_font(style, "Bitstream Charter", 10.4, INK)
        style.paragraph_format.line_spacing = 1.24
        style.paragraph_format.space_after = Pt(2.5)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBJECT
    doc.core_properties.author = AUTHOR
    doc.core_properties.keywords = SUBJECT
    doc.core_properties.comments = DESCRIPTION
    doc.save(str(REFERENCE))


def make_cover() -> None:
    width, height = 1600, 2560
    image = Image.new("RGB", (width, height), f"#{NAVY}")
    draw = ImageDraw.Draw(image)
    font_dir = Path("/usr/share/fonts/opentype/urw-base35")
    title_font = ImageFont.truetype(str(font_dir / "NimbusSansNarrow-Bold.otf"), 250)
    subtitle_font = ImageFont.truetype(str(font_dir / "NimbusSansNarrow-Regular.otf"), 78)
    author_font = ImageFont.truetype(str(font_dir / "NimbusSans-Regular.otf"), 58)
    small_font = ImageFont.truetype(str(font_dir / "NimbusSans-Regular.otf"), 31)

    draw.rectangle((0, 0, width, 42), fill=f"#{VERMILION}")
    draw.rectangle((218, 330, 242, 1940), fill=f"#{VERMILION}")
    draw.rectangle((242, 330, 270, 1180), fill=f"#{PETROL}")
    draw.text((340, 380), "FUGA", font=title_font, fill=f"#{IVORY}")
    draw.text((340, 640), "IDENTITÁRIA", font=title_font, fill=f"#{IVORY}")
    draw.multiline_text(
        (348, 1130),
        "A ANATOMIA DA\nDILUIÇÃO DO EU",
        font=subtitle_font,
        fill="#C6D5D6",
        spacing=18,
    )
    draw.line((348, 1515, 1300, 1515), fill=f"#{PETROL}", width=5)
    draw.text((348, 2080), AUTHOR.upper(), font=author_font, fill=f"#{IVORY}")
    draw.text(
        (348, 2185),
        "ENSAIO SOBRE IDENTIDADE, NARRATIVA E AUTORIA",
        font=small_font,
        fill="#AFC2C4",
    )
    image.save(COVER, format="PNG", optimize=True)
    image.save(
        COVER_JPG,
        format="JPEG",
        quality=96,
        subsampling=0,
        optimize=True,
        dpi=(300, 300),
    )


def make_fontconfig() -> None:
    FONTS.mkdir(parents=True, exist_ok=True)
    cache_dir = ROOT / "qa" / "fontconfig-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    FONTCONFIG.write_text(
        '<?xml version="1.0"?>\n'
        '<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">\n'
        "<fontconfig>\n"
        '  <include ignore_missing="yes">/etc/fonts/fonts.conf</include>\n'
        f"  <dir>{FONTS}</dir>\n"
        f"  <cachedir>{cache_dir}</cachedir>\n"
        "</fontconfig>\n",
        encoding="utf-8",
    )


def extract_figure_alts() -> list[str]:
    text = SOURCE.read_text(encoding="utf-8")
    return re.findall(r"^!\[(FIG-\d{2} — .*?)\]\(figuras/FIG-\d{2}\.png\)$", text, re.M)


def postprocess_docx() -> None:
    doc = Document(str(DOCX_OUT))
    for section in doc.sections:
        configure_section(section)

    paragraphs = list(doc.paragraphs)
    in_epilogue = False
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.startswith("EPÍLOGO —"):
            in_epilogue = True
        elif paragraph.text.startswith("MAPA DE UMA SITUAÇÃO"):
            in_epilogue = False
        if in_epilogue and paragraph.style and paragraph.style.name in {
            "Normal",
            "Body Text",
            "First Paragraph",
            "Compact",
        }:
            # Preserve the closing cadence while preventing the final sentence
            # from becoming an accidental one-line page in the technical proof.
            paragraph.paragraph_format.space_after = Pt(4.35)
        if paragraph.style and paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith("FIG-"):
            paragraph.style = get_or_add_style(doc, "Caption")
            paragraph.paragraph_format.keep_together = True
        if paragraph.style and paragraph.style.name == "Block Text":
            paragraph.paragraph_format.keep_together = True
            if (
                index + 1 < len(paragraphs)
                and paragraphs[index + 1].style
                and paragraphs[index + 1].style.name == "Block Text"
            ):
                paragraph.paragraph_format.keep_with_next = True
            # Some commands intentionally use a short quoted label followed by
            # a normal explanatory paragraph. Keep both pieces together so a
            # command never becomes an orphan at the bottom of a page.
            if "COMANDO DE AUTORIA" in paragraph.text and index + 1 < len(paragraphs):
                following = paragraphs[index + 1]
                if not (
                    following.style
                    and following.style.name.startswith("Heading")
                ):
                    paragraph.paragraph_format.keep_with_next = True
                    following.paragraph_format.keep_together = True
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True

    alts = extract_figure_alts()
    shapes = list(doc.inline_shapes)
    if len(shapes) != len(alts):
        raise RuntimeError(f"Expected {len(alts)} figures in DOCX, found {len(shapes)}")
    max_width = Cm(10.4)
    for idx, (shape, alt) in enumerate(zip(shapes, alts), 1):
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)
        doc_pr = shape._inline.docPr
        doc_pr.set("name", f"FIG-{idx:02d}")
        doc_pr.set("descr", alt)

    # Pandoc's smart typography inserts non-breaking spaces after common
    # abbreviations. The converted Charter webfont does not carry a reliable
    # NBSP glyph in LibreOffice, so those spaces render as daggers in the PDF.
    # Ordinary spaces preserve the text and eliminate the visible artifact.
    for text_node in doc.element.xpath(".//w:t"):
        if text_node.text:
            text_node.text = text_node.text.replace("\u00a0", " ")

    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBJECT
    doc.core_properties.author = AUTHOR
    doc.core_properties.last_modified_by = AUTHOR
    doc.core_properties.keywords = SUBJECT
    doc.core_properties.comments = DESCRIPTION
    doc.core_properties.created = datetime(2026, 7, 31, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 7, 31, tzinfo=timezone.utc)
    doc.save(str(DOCX_OUT))


def build_docx() -> None:
    run(
        [
            "pandoc",
            str(SOURCE),
            "--from=markdown+smart+implicit_figures",
            "--to=docx",
            f"--reference-doc={REFERENCE}",
            f"--resource-path={ROOT}",
            f"--lua-filter={SCRIPTS / 'figure_ids.lua'}",
            "--metadata=lang:pt-BR",
            f"--output={DOCX_OUT}",
        ],
        cwd=ROOT,
    )
    postprocess_docx()
    with zipfile.ZipFile(DOCX_OUT) as archive:
        damaged_member = archive.testzip()
    if damaged_member is not None:
        raise RuntimeError(f"DOCX contains a damaged member: {damaged_member}")


def render_pdf() -> None:
    render_tool = Path("/root/.codex/skills/builtins/documents/render_docx.py")
    render_dir = ROOT / "qa" / "render_docx_final"
    if render_dir.exists():
        shutil.rmtree(render_dir)
    render_dir.mkdir(parents=True, exist_ok=True)
    render_env = os.environ.copy()
    render_env["FONTCONFIG_FILE"] = str(FONTCONFIG)
    if render_tool.exists():
        run(
            [
                sys.executable,
                str(render_tool),
                str(DOCX_OUT),
                "--output_dir",
                str(render_dir),
                "--emit_pdf",
                "--dpi",
                "144",
            ],
            cwd=ROOT,
            env=render_env,
        )
        rendered_pdf = render_dir / f"{DOCX_OUT.stem}.pdf"
    else:
        run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(render_dir),
                str(DOCX_OUT),
            ],
            cwd=ROOT,
            env=render_env,
        )
        rendered_pdf = render_dir / f"{DOCX_OUT.stem}.pdf"
    if not rendered_pdf.exists():
        raise RuntimeError("PDF render did not produce an output file")
    shutil.copy2(rendered_pdf, PDF_OUT)
    # Rendering must never mutate or truncate the editable source.
    with zipfile.ZipFile(DOCX_OUT) as archive:
        damaged_member = archive.testzip()
    if damaged_member is not None:
        raise RuntimeError(f"DOCX was damaged during rendering: {damaged_member}")


def patch_epub(epub_path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="camada11-epub-") as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(epub_path) as archive:
            archive.extractall(tmp_path)

        opf_path = next(tmp_path.rglob("*.opf"))
        ET.register_namespace("", "http://www.idpf.org/2007/opf")
        ET.register_namespace("dc", "http://purl.org/dc/elements/1.1/")
        tree = ET.parse(opf_path)
        root = tree.getroot()
        ns = {
            "opf": "http://www.idpf.org/2007/opf",
            "dc": "http://purl.org/dc/elements/1.1/",
        }
        metadata = root.find("opf:metadata", ns)
        if metadata is None:
            raise RuntimeError("EPUB metadata block not found")

        date_node = metadata.find("dc:date", ns)
        if date_node is None:
            date_node = ET.SubElement(metadata, f"{{{ns['dc']}}}date")
        date_node.text = BUILD_DATE

        title_nodes = metadata.findall("dc:title", ns)
        main_title = next(
            (node for node in title_nodes if (node.text or "").strip() == TITLE),
            None,
        )
        if main_title is not None and not main_title.get("id"):
            main_title.set("id", "epub-title-main")
        subtitle_node = next(
            (node for node in title_nodes if (node.text or "").strip() == SUBTITLE),
            None,
        )
        if subtitle_node is None:
            subtitle_node = ET.SubElement(metadata, f"{{{ns['dc']}}}title")
            subtitle_node.set("id", "epub-title-subtitle")
            subtitle_node.text = SUBTITLE

        title_refinements = {
            (
                node.attrib.get("refines"),
                node.attrib.get("property"),
                (node.text or "").strip(),
            )
            for node in metadata.findall("opf:meta", ns)
        }
        if main_title is not None:
            main_refinement = (
                f"#{main_title.get('id')}",
                "title-type",
                "main",
            )
            if main_refinement not in title_refinements:
                node = ET.SubElement(metadata, f"{{{ns['opf']}}}meta")
                node.set("refines", main_refinement[0])
                node.set("property", main_refinement[1])
                node.text = main_refinement[2]
        subtitle_refinement = (
            f"#{subtitle_node.get('id')}",
            "title-type",
            "subtitle",
        )
        if subtitle_refinement not in title_refinements:
            node = ET.SubElement(metadata, f"{{{ns['opf']}}}meta")
            node.set("refines", subtitle_refinement[0])
            node.set("property", subtitle_refinement[1])
            node.text = subtitle_refinement[2]

        existing_meta = {
            (node.attrib.get("property"), (node.text or "").strip())
            for node in metadata.findall("opf:meta", ns)
        }
        access_meta = [
            ("schema:accessMode", "textual"),
            ("schema:accessMode", "visual"),
            ("schema:accessModeSufficient", "textual,visual"),
            ("schema:accessibilityFeature", "alternativeText"),
            ("schema:accessibilityFeature", "structuralNavigation"),
            ("schema:accessibilityFeature", "readingOrder"),
            ("schema:accessibilityHazard", "none"),
        ]
        for prop, value in access_meta:
            if (prop, value) not in existing_meta:
                node = ET.SubElement(metadata, f"{{{ns['opf']}}}meta")
                node.set("property", prop)
                node.text = value
        tree.write(opf_path, encoding="utf-8", xml_declaration=True)

        xhtml_ns = "http://www.w3.org/1999/xhtml"
        svg_ns = "http://www.w3.org/2000/svg"
        ET.register_namespace("", xhtml_ns)
        ET.register_namespace("svg", svg_ns)
        figures = []
        for xhtml in sorted(tmp_path.rglob("*.xhtml")):
            x_tree = ET.parse(xhtml)
            x_root = x_tree.getroot()
            changed = False
            if xhtml.name == "cover.xhtml":
                for image_node in x_root.findall(f".//{{{svg_ns}}}image"):
                    image_node.set("role", "presentation")
                    image_node.set("aria-hidden", "true")
                    changed = True
            for fig in x_root.findall(f".//{{{xhtml_ns}}}figure"):
                figures.append((xhtml, fig))
                number = len(figures)
                fig.set("id", f"fig-{number:02d}")
                img = fig.find(f".//{{{xhtml_ns}}}img")
                if img is None:
                    raise RuntimeError(f"Figure without image in {xhtml}")
                alt = (img.get("alt") or "").strip()
                expected = f"FIG-{number:02d}"
                if not alt.startswith(expected):
                    raise RuntimeError(
                        f"Unexpected alt text for figure {number}: {alt!r}"
                    )
                caption = fig.find(f"{{{xhtml_ns}}}figcaption")
                if caption is None:
                    caption = ET.SubElement(fig, f"{{{xhtml_ns}}}figcaption")
                    caption.text = alt
                changed = True
            if changed:
                x_tree.write(xhtml, encoding="utf-8", xml_declaration=True)
        if len(figures) != 21:
            raise RuntimeError(f"Expected 21 semantic figures in EPUB, found {len(figures)}")

        rebuilt = epub_path.with_suffix(".rebuilt.epub")
        with zipfile.ZipFile(rebuilt, "w") as archive:
            mimetype = tmp_path / "mimetype"
            archive.write(mimetype, "mimetype", compress_type=zipfile.ZIP_STORED)
            for path in sorted(tmp_path.rglob("*")):
                if not path.is_file() or path == mimetype:
                    continue
                archive.write(
                    path,
                    path.relative_to(tmp_path).as_posix(),
                    compress_type=zipfile.ZIP_DEFLATED,
                )
        rebuilt.replace(epub_path)


def build_epub() -> None:
    run(
        [
            "pandoc",
            str(SOURCE),
            "--from=markdown+smart+implicit_figures",
            "--to=epub3",
            "--split-level=1",
            "--toc-depth=1",
            "--epub-title-page=true",
            f"--epub-cover-image={COVER}",
            f"--css={SCRIPTS / 'epub.css'}",
            f"--resource-path={ROOT}",
            f"--lua-filter={SCRIPTS / 'figure_ids.lua'}",
            "--metadata=lang:pt-BR",
            f"--output={EPUB_OUT}",
        ],
        cwd=ROOT,
    )
    patch_epub(EPUB_OUT)
    if not EPUB_OUT.exists() or EPUB_OUT.stat().st_size < 100_000:
        raise RuntimeError("EPUB output is missing or unexpectedly small")
    with zipfile.ZipFile(EPUB_OUT) as archive:
        damaged_member = archive.testzip()
        if damaged_member is not None:
            raise RuntimeError(f"EPUB contains a damaged member: {damaged_member}")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    make_cover()
    make_fontconfig()
    make_reference_docx()
    build_docx()
    render_pdf()
    build_epub()
    print(f"Built: {DOCX_OUT}")
    print(f"Built: {PDF_OUT}")
    print(f"Built: {EPUB_OUT}")
    print(f"Built: {COVER_JPG}")


if __name__ == "__main__":
    main()
