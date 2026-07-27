#!/usr/bin/env python3
"""Build and style the editorial map DOCX from its canonical Markdown source."""

from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "1F3A4D"
TEAL = "1F6D6A"
PALE_TEAL = "E8F2F1"
PALE_GRAY = "F3F5F6"
MID_GRAY = "6B7280"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = f"w:{edge}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def set_run_font(run, name: str, size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(docx_path: Path) -> None:
    doc = Document(docx_path)
    doc.core_properties.title = "Fuga Identitária — Mapa Mestre de Produção 2.0"
    doc.core_properties.author = "Sol Lima"
    doc.core_properties.subject = "Estrutura, produção, revisão e preparação para Kindle Create"
    doc.core_properties.keywords = "fuga identitária, autoria, narrativa, livro, Kindle Create"
    doc.core_properties.comments = "Documento editorial de trabalho."

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.85)
        section.right_margin = Cm(1.85)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    for style_name in ("Body Text", "Body Text 2", "Body Text 3"):
        if style_name in styles:
            st = styles[style_name]
            st.font.name = "Liberation Serif"
            st._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")
            st.font.size = Pt(10.5)

    heading_specs = {
        "Title": (28, NAVY, 0, 10),
        "Subtitle": (17, TEAL, 0, 12),
        "Heading 1": (22, NAVY, 20, 9),
        "Heading 2": (16, TEAL, 16, 6),
        "Heading 3": (12.5, NAVY, 12, 4),
        "Heading 4": (11, TEAL, 9, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        if name not in styles:
            continue
        st = styles[name]
        st.font.name = "Liberation Sans"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.widow_control = True

    if "Block Text" in styles:
        block = styles["Block Text"]
        block.font.name = "Liberation Serif"
        block._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")
        block.font.size = Pt(10.5)
        block.font.italic = True
        block.font.color.rgb = RGBColor.from_string(NAVY)
        block.paragraph_format.left_indent = Cm(0.6)
        block.paragraph_format.right_indent = Cm(0.4)
        block.paragraph_format.space_before = Pt(6)
        block.paragraph_format.space_after = Pt(8)

    # Cover styling and intentional page break.
    cover_date_found = False
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if idx == 0 and text == "Fuga Identitária":
            paragraph.style = styles["Title"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(78)
        elif text == "Mapa Mestre de Estrutura, Produção e Revisão — Versão 2.0":
            paragraph.style = styles["Subtitle"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif idx < 12 and text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "Liberation Sans", 10, MID_GRAY)
        if "26 de julho de 2026" in text:
            paragraph.paragraph_format.space_after = Pt(18)
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            cover_date_found = True
            break

    # Keep list items and short paragraphs together where practical.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if text.startswith(">"):
            paragraph.paragraph_format.left_indent = Cm(0.6)
        if text.startswith("**") and text.endswith("**"):
            paragraph.paragraph_format.keep_with_next = True

    # Style tables for legibility in print and PDF.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    set_cell_shading(cell, NAVY)
                elif row_index % 2 == 0:
                    set_cell_shading(cell, PALE_GRAY)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2.5)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, "Liberation Sans", 8.6, WHITE if row_index == 0 else None)
                        if row_index == 0:
                            run.font.bold = True

    # Headers and footers.
    for section_index, section in enumerate(doc.sections):
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = "FUGA IDENTITÁRIA  •  MAPA MESTRE 2.0"
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_para.paragraph_format.space_after = Pt(0)
        for run in header_para.runs:
            set_run_font(run, "Liberation Sans", 8, MID_GRAY)

        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.paragraph_format.space_before = Pt(0)
        label = footer_para.add_run("Sol Lima  •  ")
        set_run_font(label, "Liberation Sans", 8, MID_GRAY)
        add_page_number(footer_para)
        for run in footer_para.runs:
            set_run_font(run, "Liberation Sans", 8, MID_GRAY)

    # Cover should not display header/footer.
    if doc.sections:
        doc.sections[0].different_first_page_header_footer = True

    if not cover_date_found:
        raise RuntimeError("Cover date paragraph not found; refusing to create an unpaginated cover.")

    doc.save(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            "pandoc",
            str(args.source),
            "--from=gfm",
            "--to=docx",
            "--output",
            str(args.output),
        ],
        check=True,
    )
    style_document(args.output)


if __name__ == "__main__":
    main()
