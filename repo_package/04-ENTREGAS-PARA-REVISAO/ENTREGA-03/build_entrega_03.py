from __future__ import annotations

import base64
import hashlib
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Iterable

import cairosvg
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD = ROOT / "FUGA_IDENTITARIA_ENTREGA_03_TEXTO.md"
FIGDIR = ROOT / "figuras"
DOCX = ROOT / "Fuga_Identitaria_Entrega_03_Parte_III.docx"
PDF = ROOT / "Fuga_Identitaria_Entrega_03_Parte_III.pdf"

NAVY = "0B1020"
VERMILION = "D84A3A"
IVORY = "F4EFE6"
PETROLEUM = "5F7F83"
INK = "20242B"
LIGHT = "FBF8F2"

ALT = {
    "FIG-07": "Diagrama dos primeiros mapas do eu. Família e cuidadores, pares e escola, cultura e fé, telas e plataformas, e acontecimentos e condições enviam interpretações iniciais ao eu. O leitor pode examinar, integrar, revisar ou recusar esses mapas.",
    "FIG-08": "Ciclo de recomendação. Atenção observada gera padrão inferido, que produz conteúdo recomendado, nova atenção e atualização do modelo. Uma nota ressalta que atenção não equivale a intenção ou identidade e que existe assimetria entre dados da plataforma e visão do usuário.",
    "FIG-09": "Diagrama de quatro dimensões: eu vivido, eu ideal, eu exibido e eu percebido pela audiência. As setas mostram tensões e traduções entre as dimensões, sem tratar a discrepância como diagnóstico.",
    "FIG-10": "Comparação entre causa que serve à pessoa e causa que passa a usar a pessoa. No caminho saudável, injustiça leva a linguagem, solidariedade, ação e revisão. No caminho de captura, vocabulário exclusivo, inimigo, pureza, punição à dúvida e custo de saída crescente reduzem autoria.",
}


def svg_wrap(content: str, title: str, desc: str, w: int = 2000, h: int = 1250) -> str:
    return f'''<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}" role="img" aria-labelledby="title desc">
<title id="title">{title}</title><desc id="desc">{desc}</desc>
<rect width="{w}" height="{h}" fill="#{IVORY}"/>
<style>
.title{{font-family:'Nimbus Sans Narrow','Arial Narrow',sans-serif;font-size:70px;font-weight:700;fill:#{NAVY};letter-spacing:1px}}
.sub{{font-family:'Nimbus Sans',Arial,sans-serif;font-size:32px;fill:#{INK}}}
.head{{font-family:'Nimbus Sans',Arial,sans-serif;font-size:36px;font-weight:700;fill:#{NAVY}}}
.body{{font-family:'Nimbus Sans',Arial,sans-serif;font-size:29px;fill:#{INK}}}
.small{{font-family:'Nimbus Sans',Arial,sans-serif;font-size:25px;fill:#{INK}}}
.box{{fill:#{LIGHT};stroke:#{PETROLEUM};stroke-width:5;rx:26}}
.arrow{{stroke:#{VERMILION};stroke-width:8;fill:none;marker-end:url(#arrow)}}
.soft{{stroke:#{PETROLEUM};stroke-width:5;fill:none;marker-end:url(#arrow2)}}
</style>
<defs><marker id="arrow" markerWidth="14" markerHeight="14" refX="11" refY="5" orient="auto"><path d="M0,0 L0,10 L12,5 z" fill="#{VERMILION}"/></marker><marker id="arrow2" markerWidth="14" markerHeight="14" refX="11" refY="5" orient="auto"><path d="M0,0 L0,10 L12,5 z" fill="#{PETROLEUM}"/></marker></defs>
{content}</svg>'''


def make_figures() -> None:
    FIGDIR.mkdir(parents=True, exist_ok=True)

    content = f'''
<text x="1000" y="100" text-anchor="middle" class="title">DE ONDE CHEGAM OS MAPAS DO EU</text>
<text x="1000" y="150" text-anchor="middle" class="sub">Primeiras interpretações não são destino: podem ser examinadas</text>
<rect x="735" y="465" width="530" height="190" class="box" fill="#{NAVY}"/>
<text x="1000" y="535" text-anchor="middle" class="head">MAPAS DO EU</text>
<text x="1000" y="585" text-anchor="middle" class="body">amor · perigo · valor</text>
<text x="1000" y="625" text-anchor="middle" class="body">voz · pertencimento</text>
<rect x="80" y="240" width="420" height="150" class="box"/><text x="290" y="300" text-anchor="middle" class="head">Família e cuidadores</text><text x="290" y="345" text-anchor="middle" class="small">presença · limite · reação</text>
<rect x="790" y="225" width="420" height="150" class="box"/><text x="1000" y="285" text-anchor="middle" class="head">Pares e escola</text><text x="1000" y="330" text-anchor="middle" class="small">comparação · regra · status</text>
<rect x="1500" y="240" width="420" height="150" class="box"/><text x="1710" y="300" text-anchor="middle" class="head">Cultura e fé</text><text x="1710" y="345" text-anchor="middle" class="small">símbolo · memória · valor</text>
<rect x="210" y="740" width="470" height="150" class="box"/><text x="445" y="800" text-anchor="middle" class="head">Telas e plataformas</text><text x="445" y="845" text-anchor="middle" class="small">repetição · recomendação</text>
<rect x="1320" y="740" width="470" height="150" class="box"/><text x="1555" y="800" text-anchor="middle" class="head">Acontecimentos</text><text x="1555" y="845" text-anchor="middle" class="small">perdas · condições · encontros</text>
<path d="M500 350 C650 380 720 430 790 485" class="soft"/><path d="M1000 375 L1000 450" class="soft"/><path d="M1500 350 C1350 380 1280 430 1210 485" class="soft"/><path d="M680 790 C790 720 820 680 850 650" class="soft"/><path d="M1320 790 C1210 720 1180 680 1150 650" class="soft"/>
<path d="M1000 655 L1000 970" class="arrow"/>
<rect x="455" y="970" width="1090" height="170" rx="28" fill="#{NAVY}"/>
<text x="1000" y="1035" text-anchor="middle" style="font-family:'Nimbus Sans',Arial;font-size:38px;font-weight:700;fill:#{IVORY}">AUTORIA ADULTA</text>
<text x="1000" y="1090" text-anchor="middle" style="font-family:'Nimbus Sans',Arial;font-size:31px;fill:#{IVORY}">examinar · integrar · revisar · recusar</text>
'''
    write_svg_png("FIG-07_DE_ONDE_CHEGAM_OS_MAPAS_DO_EU", svg_wrap(content, "De onde chegam os mapas do eu", ALT["FIG-07"]))

    boxes = [(100,410,"ATENÇÃO\nOBSERVADA"),(480,230,"PADRÃO\nINFERIDO"),(980,230,"CONTEÚDO\nRECOMENDADO"),(1480,410,"NOVA\nATENÇÃO"),(980,720,"MODELO\nATUALIZADO"),(480,720,"PERFIL DE\nPREFERÊNCIA")]
    content = f'<text x="1000" y="100" text-anchor="middle" class="title">O CICLO DA RECOMENDAÇÃO</text><text x="1000" y="150" text-anchor="middle" class="sub">A resposta futura também depende do que foi tornado visível</text>'
    for x,y,label in boxes:
        content += f'<rect x="{x}" y="{y}" width="350" height="150" class="box"/>'
        lines=label.split('\n')
        content += ''.join(f'<text x="{x+175}" y="{y+62+i*42}" text-anchor="middle" class="head">{ln}</text>' for i,ln in enumerate(lines))
    pts=[((450,485),(520,360)),((830,305),(980,305)),((1330,305),(1510,420)),((1655,560),(1330,760)),((980,795),(830,795)),((480,760),(260,560))]
    for (x1,y1),(x2,y2) in pts: content+=f'<path d="M{x1} {y1} L{x2} {y2}" class="arrow"/>'
    content += f'''
<rect x="335" y="990" width="1330" height="150" rx="28" fill="#{NAVY}"/>
<text x="1000" y="1045" text-anchor="middle" style="font-family:'Nimbus Sans';font-size:34px;font-weight:700;fill:#{IVORY}">ATENÇÃO ≠ INTENÇÃO ≠ IDENTIDADE</text>
<text x="1000" y="1092" text-anchor="middle" style="font-family:'Nimbus Sans';font-size:28px;fill:#{IVORY}">usuário vê a seleção · plataforma mede sinais em escala</text>'''
    write_svg_png("FIG-08_O_CICLO_DA_RECOMENDACAO", svg_wrap(content, "O ciclo da recomendação", ALT["FIG-08"]))

    content=f'''
<text x="1000" y="100" text-anchor="middle" class="title">EU VIVIDO E EU EXIBIDO</text>
<text x="1000" y="150" text-anchor="middle" class="sub">Diferença não é falsidade; o risco é o personagem proibir a pessoa</text>
<circle cx="600" cy="455" r="235" fill="#{LIGHT}" stroke="#{PETROLEUM}" stroke-width="6"/><text x="600" y="420" text-anchor="middle" class="head">EU VIVIDO</text><text x="600" y="470" text-anchor="middle" class="body">experiência · limite</text><text x="600" y="510" text-anchor="middle" class="body">ambivalência · intimidade</text>
<circle cx="1400" cy="455" r="235" fill="#{LIGHT}" stroke="#{VERMILION}" stroke-width="6"/><text x="1400" y="420" text-anchor="middle" class="head">EU IDEAL</text><text x="1400" y="470" text-anchor="middle" class="body">aspiração · dever</text><text x="1400" y="510" text-anchor="middle" class="body">medo · promessa</text>
<circle cx="600" cy="895" r="235" fill="#{LIGHT}" stroke="#{VERMILION}" stroke-width="6"/><text x="600" y="860" text-anchor="middle" class="head">EU EXIBIDO</text><text x="600" y="910" text-anchor="middle" class="body">recorte · estética</text><text x="600" y="950" text-anchor="middle" class="body">legenda · marca</text>
<circle cx="1400" cy="895" r="235" fill="#{LIGHT}" stroke="#{PETROLEUM}" stroke-width="6"/><text x="1400" y="850" text-anchor="middle" class="head">EU PERCEBIDO</text><text x="1400" y="895" text-anchor="middle" class="head">PELA AUDIÊNCIA</text><text x="1400" y="945" text-anchor="middle" class="body">expectativa · projeção</text>
<path d="M835 455 L1165 455" class="soft"/><path d="M600 690 L600 660" class="soft"/><path d="M835 895 L1165 895" class="soft"/><path d="M1400 690 L1400 660" class="soft"/>
<rect x="810" y="580" width="380" height="190" rx="95" fill="#{NAVY}"/><text x="1000" y="650" text-anchor="middle" style="font-family:'Nimbus Sans';font-size:33px;font-weight:700;fill:#{IVORY}">TENSÃO</text><text x="1000" y="700" text-anchor="middle" style="font-family:'Nimbus Sans';font-size:28px;fill:#{IVORY}">não é diagnóstico</text><text x="1000" y="742" text-anchor="middle" style="font-family:'Nimbus Sans';font-size:25px;fill:#{IVORY}">é pista para exame</text>
'''
    write_svg_png("FIG-09_EU_VIVIDO_E_EU_EXIBIDO", svg_wrap(content, "Eu vivido e eu exibido", ALT["FIG-09"]))

    content=f'''
<text x="1000" y="100" text-anchor="middle" class="title">QUANDO A CAUSA COMEÇA A USAR A PESSOA</text>
<text x="1000" y="150" text-anchor="middle" class="sub">Comparar mecanismos não iguala histórias, ideias ou danos</text>
<rect x="80" y="225" width="840" height="900" rx="30" fill="#{LIGHT}" stroke="#{PETROLEUM}" stroke-width="6"/>
<text x="500" y="295" text-anchor="middle" class="head">A CAUSA SERVE À PESSOA</text>
<rect x="1080" y="225" width="840" height="900" rx="30" fill="#{LIGHT}" stroke="#{VERMILION}" stroke-width="6"/>
<text x="1500" y="295" text-anchor="middle" class="head">A PESSOA SERVE À CAUSA</text>
'''
    left=["Injustiça ou necessidade","Linguagem e memória","Solidariedade","Ação responsável","Revisão · crítica · saída"]
    right=["Vocabulário exclusivo","Inimigo indispensável","Teste de pureza","Dúvida punida","Missão ocupa a vida","Custo de saída cresce"]
    for i,s in enumerate(left):
        y=350+i*145; content+=f'<rect x="190" y="{y}" width="620" height="95" class="box"/><text x="500" y="{y+59}" text-anchor="middle" class="body">{s}</text>'
        if i<len(left)-1: content+=f'<path d="M500 {y+95} L500 {y+135}" class="soft"/>'
    for i,s in enumerate(right):
        y=335+i*125; content+=f'<rect x="1190" y="{y}" width="620" height="85" class="box"/><text x="1500" y="{y+54}" text-anchor="middle" class="body">{s}</text>'
        if i<len(right)-1: content+=f'<path d="M1745 {y+85} L1745 {y+115}" class="arrow"/>'
    content += f'<text x="500" y="1080" text-anchor="middle" class="head">autoria ampliada</text><text x="1500" y="1080" text-anchor="middle" class="head">autoria terceirizada</text>'
    write_svg_png("FIG-10_QUANDO_A_CAUSA_COMECA_A_USAR_A_PESSOA", svg_wrap(content, "Quando a causa começa a usar a pessoa", ALT["FIG-10"]))


def write_svg_png(stem: str, svg: str) -> None:
    svg_path = FIGDIR / f"{stem}.svg"
    png_path = FIGDIR / f"{stem}.png"
    svg_path.write_text(svg, encoding="utf-8")
    cairosvg.svg2png(bytestring=svg.encode("utf-8"), write_to=str(png_path), output_width=2000, output_height=1250)


def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def set_paragraph_shading(p, fill: str) -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pPr.append(shd)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])


def set_page_number_start(section, start: int) -> None:
    sectPr = section._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sectPr.append(pg)
    pg.set(qn('w:start'), str(start))


def set_image_alt(inline_shape, title: str, description: str) -> None:
    docPr = inline_shape._inline.docPr
    docPr.set('title', title)
    docPr.set('descr', description)


def add_formatted_run(p, text: str) -> None:
    # inline **bold** and *italic* parser
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('*') and part.endswith('*'):
            r=p.add_run(part[1:-1]); r.italic=True
        else:
            p.add_run(part)


def add_body_paragraph(doc: Document, text: str, style: str = 'Body Text'):
    p=doc.add_paragraph(style=style)
    add_formatted_run(p,text)
    return p


def define_styles(doc: Document) -> None:
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Bitstream Charter'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Bitstream Charter'); normal.font.size=Pt(10.4); normal.font.color.rgb=RGBColor.from_string(INK)
    pf=normal.paragraph_format; pf.space_after=Pt(4); pf.line_spacing=1.3; pf.alignment=WD_ALIGN_PARAGRAPH.LEFT

    for name,size,before,after in [('Title',29,0,10),('Subtitle',15,0,12),('Heading 1',22,0,10),('Heading 2',15,12,6),('Heading 3',12,8,4)]:
        st=styles[name]; st.font.name='Nimbus Sans Narrow'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Nimbus Sans Narrow'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    styles['Heading 1'].paragraph_format.page_break_before=False
    styles['Body Text'].font.name='Bitstream Charter'; styles['Body Text']._element.rPr.rFonts.set(qn('w:eastAsia'),'Bitstream Charter'); styles['Body Text'].font.size=Pt(10.4); styles['Body Text'].paragraph_format.space_after=Pt(4); styles['Body Text'].paragraph_format.line_spacing=1.3
    styles['Quote'].font.name='Bitstream Charter'; styles['Quote']._element.rPr.rFonts.set(qn('w:eastAsia'),'Bitstream Charter'); styles['Quote'].font.size=Pt(10.4); styles['Quote'].font.italic=False; styles['Quote'].paragraph_format.left_indent=Cm(.65); styles['Quote'].paragraph_format.right_indent=Cm(.35); styles['Quote'].paragraph_format.space_before=Pt(5); styles['Quote'].paragraph_format.space_after=Pt(5); styles['Quote'].paragraph_format.line_spacing=1.2
    if 'Figure Caption' not in styles:
        st=styles.add_style('Figure Caption',WD_STYLE_TYPE.PARAGRAPH)
    st=styles['Figure Caption']; st.font.name='Nimbus Sans'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Nimbus Sans'); st.font.size=Pt(8.5); st.font.italic=True; st.font.color.rgb=RGBColor.from_string(INK); st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_after=Pt(9); st.paragraph_format.keep_with_next=False
    if 'Reference' not in styles:
        st=styles.add_style('Reference',WD_STYLE_TYPE.PARAGRAPH)
    st=styles['Reference']; st.font.name='Bitstream Charter'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Bitstream Charter'); st.font.size=Pt(8.4); st.paragraph_format.left_indent=Cm(.45); st.paragraph_format.first_line_indent=Cm(-.45); st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.12
    if 'Command' not in styles:
        st=styles.add_style('Command',WD_STYLE_TYPE.PARAGRAPH)
    st=styles['Command']; st.font.name='Nimbus Sans'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Nimbus Sans'); st.font.size=Pt(9.2); st.paragraph_format.left_indent=Cm(.55); st.paragraph_format.right_indent=Cm(.35); st.paragraph_format.space_after=Pt(2.5); st.paragraph_format.line_spacing=1.15


def configure_section(section) -> None:
    section.page_width=Cm(14); section.page_height=Cm(21)
    section.top_margin=Cm(1.6); section.bottom_margin=Cm(1.75); section.left_margin=Cm(1.8); section.right_margin=Cm(1.5)
    section.header_distance=Cm(.7); section.footer_distance=Cm(.75)
    sectPr=section._sectPr
    mirror=OxmlElement('w:mirrorMargins'); sectPr.append(mirror)


def make_docx() -> None:
    doc=Document(); define_styles(doc)
    for sec in doc.sections: configure_section(sec)
    sec0=doc.sections[0]
    # technical sheet
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(0)
    r=p.add_run('FUGA IDENTITÁRIA'); r.font.name='Nimbus Sans Narrow'; r.font.size=Pt(27); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('A Anatomia da Diluição do Eu'); r.font.name='Nimbus Sans Narrow'; r.font.size=Pt(16); r.italic=True; r.font.color.rgb=RGBColor.from_string(PETROLEUM)
    doc.add_paragraph('')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('PARTE III — OS FORNECEDORES DO EU'); r.font.name='Nimbus Sans Narrow'; r.font.size=Pt(17.5); r.bold=True; r.font.color.rgb=RGBColor.from_string(VERMILION)
    doc.add_paragraph('')
    for label,value in [('Autora','Sol Lima'),('Entrega','3 do ciclo numerado de 0 a 7'),('Estado editorial','Escrita integral → Auditoria factual → Revisão editorial → Revisão de Sol'),('Faixa editorial','início na página 114'),('Uso','cópia de revisão; ainda não aprovada nem fundida ao manuscrito mestre')]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        a=p.add_run(label+': '); a.bold=True; a.font.name='Nimbus Sans'; b=p.add_run(value); b.font.name='Nimbus Sans';
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); r=p.add_run('27 de julho de 2026'); r.font.name='Nimbus Sans'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(PETROLEUM)

    sec=doc.add_section(WD_SECTION.NEW_PAGE); configure_section(sec); set_page_number_start(sec,114); sec.footer.is_linked_to_previous=False; add_page_field(sec.footer.paragraphs[0]); sec.header.is_linked_to_previous=False

    text=MD.read_text(encoding='utf-8')
    # start at second Part III heading
    marker='# Parte III — Os fornecedores do eu'
    start=text.find(marker, text.find(marker)+1)
    lines=text[start:].splitlines()
    in_refs=False; quote_lines=[]
    fig_num=7

    def flush_quote():
        nonlocal quote_lines
        if not quote_lines: return
        for i,q in enumerate(quote_lines):
            p=doc.add_paragraph(style='Command' if any('**' in z for z in quote_lines) else 'Quote')
            set_paragraph_shading(p, 'F6F0E7')
            pPr=p._p.get_or_add_pPr(); borders=pPr.find(qn('w:pBdr'))
            if borders is None:
                borders=OxmlElement('w:pBdr'); pPr.append(borders)
            left=OxmlElement('w:left'); left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'8'); left.set(qn('w:color'),VERMILION); borders.append(left)
            add_formatted_run(p,q)
        quote_lines=[]

    first_h1=True
    for raw in lines:
        line=raw.rstrip()
        if line.startswith('>'):
            quote_lines.append(line[1:].lstrip())
            continue
        flush_quote()
        if not line or line=='---':
            continue
        if line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            if m:
                caption,path=m.groups(); img=ROOT/path
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2)
                shape=p.add_run().add_picture(str(img),width=Cm(10.6)); key=f'FIG-{fig_num:02d}'; set_image_alt(shape,caption,ALT[key]);
                cp=doc.add_paragraph(caption,style='Figure Caption');
                fig_num+=1
            continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Heading 1')
            if not first_h1:
                p.paragraph_format.page_break_before=True
            add_formatted_run(p,line[2:]); first_h1=False
            if line.startswith('# Notas de pesquisa'): in_refs=True
            continue
        if line.startswith('## '):
            p=doc.add_paragraph(style='Heading 2'); add_formatted_run(p,line[3:]); continue
        if line.startswith('### '):
            p=doc.add_paragraph(style='Heading 3'); add_formatted_run(p,line[4:]); continue
        if re.match(r'^[-*] ',line):
            p=doc.add_paragraph(style='List Bullet'); add_formatted_run(p,line[2:]); continue
        if re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(style='List Number'); add_formatted_run(p,re.sub(r'^\d+\. ','',line)); continue
        p=add_body_paragraph(doc,line,'Reference' if in_refs else 'Body Text')
        if in_refs:
            p.paragraph_format.keep_together=False
    flush_quote()

    props=doc.core_properties; props.title='Fuga Identitária — Entrega 3 — Parte III'; props.author='Sol Lima'; props.subject='Os fornecedores do eu'; props.keywords='identidade, autoria, família, algoritmos, influenciadores, movimentos, rótulos'; props.comments='Cópia de revisão. Ainda não aprovada nem fundida ao manuscrito mestre.'
    doc.save(DOCX)


def convert_pdf() -> None:
    out=ROOT/'_lo_pdf'; out.mkdir(exist_ok=True)
    env=os.environ.copy(); env['HOME']=str(ROOT/'_lo_home'); Path(env['HOME']).mkdir(exist_ok=True)
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(out),str(DOCX)],check=True,env=env,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    made=out/(DOCX.stem+'.pdf')
    shutil.copy2(made,PDF)


def main() -> None:
    make_figures(); make_docx(); convert_pdf()
    print(DOCX); print(PDF)

if __name__=='__main__': main()
